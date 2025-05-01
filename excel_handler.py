import openpyxl
from models import Student

class ExcelHandler:
    def __init__(self, file_path):
        self.filepath = file_path

    def get_excel_file_data(self):

        # Loading the workbook
        wb = openpyxl.load_workbook(self.filepath)
        sheet = wb.active

        # Initialising list of students list
        list_of_students = []

        # Iterate over rows starting from 4th row
        for row in sheet.iter_rows(min_row=4, max_col=21, values_only=True):  # 21 corresponds to column U
            name = row[0]  # Name from column A
            class_year = row[1]  # Class from column B

            # Check if the class year is empty
            if not class_year:
                break  # Stop if the class year is empty
            
            # Initialising list of classes variable
            list_of_classes = []

            # Looping through all the classes from column 4-21 for each student
            for class_name in row[3:21]:
                if class_name:
                    # Appending class name to list of classes
                    list_of_classes.append(class_name)
         
            # Creating a student class object for each student
            student = Student(student_name=name, classes=list_of_classes)

            # Appending to a list of student class objects
            list_of_students.append(student)

        # Closing the workbook
        wb.close()

        return list_of_students


    def process_file(self, list_of_students):

        # Create a new workbook
        wb = openpyxl.Workbook()

        # Select the active worksheet
        ws = wb.active

        # Add headers
        ws['A1'] = 'Name'
        ws['B1'] = 'Class'

        # Initialise the row starting point
        row_index = 2
        
        for student in list_of_students:
            print(f'Creating entry for student: {student.student_name}')

            # Insert student name for each class
            for class_name in student.classes:
                # Insert student name
                ws.cell(row=row_index, column=1, value=student.student_name)
                
                # Insert class
                ws.cell(row=row_index, column=2, value=class_name)
                
                # Move to the next row
                row_index += 1

                # Save the workbook with a desired filename
                wb.save("processed_file.xlsx")

        wb.close()

